"""Read Jotform membership export (.xlsx) and print one JSON array to stdout."""
from __future__ import annotations

import json
import sys
from datetime import datetime
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    print("Install openpyxl: pip install openpyxl", file=sys.stderr)
    sys.exit(1)

COL = {
    "submission_date": 0,
    "first_name": 1,
    "last_name": 2,
    "membership_type": 5,
    "membership_duration": 6,
    "status": 7,
    "email": 9,
    "additional_adult": 19,
    "additional_first": 20,
    "additional_last": 21,
    "junior_1": 23,
    "junior_2": 25,
    "junior_3": 27,
}


def cell(row, index):
    if index >= len(row):
        return None
    value = row[index]
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def parse_submission_date(raw):
    if raw is None:
        return None
    if isinstance(raw, datetime):
        return raw.date().isoformat()
    text = str(raw).strip()
    if not text:
        return None
    for fmt in ("%b %d, %Y", "%B %d, %Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            continue
    return None


def main():
    if len(sys.argv) < 2:
        print("Usage: python jotform_xlsx_to_json.py <path.xlsx>", file=sys.stderr)
        sys.exit(1)

    path = Path(sys.argv[1])
    if not path.is_file():
        print(f"File not found: {path}", file=sys.stderr)
        sys.exit(1)

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows_out = []

    for row in ws.iter_rows(min_row=2, values_only=True):
        row = list(row) if row else []
        email = cell(row, COL["email"])
        if not email:
            continue
        juniors = [
            cell(row, COL["junior_1"]),
            cell(row, COL["junior_2"]),
            cell(row, COL["junior_3"]),
        ]
        juniors = [j for j in juniors if j]

        rows_out.append(
            {
                "email": email.lower(),
                "submission_date": parse_submission_date(cell(row, COL["submission_date"])),
                "primary_first_name": cell(row, COL["first_name"]) or "",
                "primary_last_name": cell(row, COL["last_name"]) or "",
                "membership_type_raw": cell(row, COL["membership_type"]) or "",
                "membership_duration_raw": cell(row, COL["membership_duration"]) or "",
                "status_raw": cell(row, COL["status"]) or "",
                "additional_adult": (cell(row, COL["additional_adult"]) or "").lower()
                in ("yes", "y", "true", "1"),
                "additional_first_name": cell(row, COL["additional_first"]) or "",
                "additional_last_name": cell(row, COL["additional_last"]) or "",
                "junior_names": juniors,
            }
        )

    json.dump(rows_out, sys.stdout)


if __name__ == "__main__":
    main()
